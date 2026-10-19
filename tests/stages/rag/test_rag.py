"""문서 검색(RAG): 파서 · 청킹 · BM25 · 융합 · API 전체 흐름."""

import io
import struct
import zipfile

import pytest

from app.core.korean import tokenize
from app.stages.rag import chunking, parsers
from app.stages.rag.bm25 import BM25Index
from app.stages.rag.fusion import rrf, weighted
from tests.helpers import DOCS, LAW, model_id, wait_job


def test_parse_text_encodings_and_html():
    assert parsers.parse("a.txt", "안녕하세요 문서".encode("cp949")).pages == ["안녕하세요 문서"]
    doc = parsers.parse("b.html", "<html><title>제목</title><script>x()</script><p>본문 <b>굵게</b></p></html>".encode())
    assert doc.title == "제목" and doc.pages == ["본문 굵게"]
    with pytest.raises(parsers.ParseError):
        parsers.parse("c.exe", b"MZ")
    with pytest.raises(parsers.ParseError):
        parsers.parse("d.txt", b"")


def test_parse_docx_keeps_tables_in_order():
    from docx import Document as Docx

    d = Docx()
    d.add_paragraph("첫 문단")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "항목", "값"
    d.add_paragraph("마지막 문단")
    buf = io.BytesIO()
    d.save(buf)
    doc = parsers.parse("x.docx", buf.getvalue())
    assert doc.pages[0].splitlines() == ["첫 문단", "항목 | 값", "마지막 문단"]


def test_parse_hwpx():
    ns = 'xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph"'
    sec = f'<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" {ns}><hp:p><hp:run><hp:t>제1조(목적) 테스트</hp:t></hp:run></hp:p><hp:p><hp:run><hp:t>둘째</hp:t><hp:t> 문단</hp:t></hp:run></hp:p></hs:sec>'
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("mimetype", "application/hwp+zip")
        z.writestr("Contents/section0.xml", sec)
        z.writestr("Contents/section1.xml", sec.replace("둘째", "셋째"))
    doc = parsers.parse("x.hwpx", buf.getvalue())
    assert doc.page_kind == "section" and len(doc.pages) == 2
    assert doc.pages[0] == "제1조(목적) 테스트\n둘째 문단"


def test_hwp_record_text_extraction():
    # PARA_TEXT 레코드: '가' + 탭(인라인 제어, 8글자) + '나' + 문단 끝(13)
    payload = struct.pack("<H", ord("가")) + struct.pack("<8H", 9, 0, 0, 0, 0, 0, 0, 9) + struct.pack("<2H", ord("나"), 13)
    record = struct.pack("<I", parsers.HWPTAG_PARA_TEXT | (len(payload) << 20)) + payload
    other = struct.pack("<I", 66 | (4 << 20)) + b"\x00" * 4  # 다른 레코드는 건너뜀
    assert parsers._hwp_section_text(other + record) == "가\t나"


def test_legal_chunking_headings():
    chunks = chunking.chunk_pages([LAW], {"method": "legal", "size": 100, "overlap": 0, "min_chars": 5}, paged=False)
    headings = [c.meta.get("heading") for c in chunks]
    assert "제1조(목적)" in headings and "제14조의2(공공언어 개선)" in headings and "【주문】" in headings
    clause2 = next(c for c in chunks if c.text.startswith("②"))
    assert clause2.meta["heading"] == "제14조(공문서의 작성)"  # 항 단위로 나뉘어도 조문 제목 유지
    assert chunking.embed_text(clause2.text, "국어기본법", clause2.meta["heading"], {"include_title": True}).startswith(
        "국어기본법 · 제14조(공문서의 작성)\n②"
    )


@pytest.mark.parametrize("method", ["fixed", "recursive", "sentence", "page"])
def test_chunk_size_and_coverage(method):
    text = ("공공기관의 공문서는 한글로 작성한다. " * 40 + "\n\n") * 3
    chunks = chunking.chunk_pages([text, "둘째 쪽 내용입니다. " * 5], {"method": method, "size": 300, "overlap": 50})
    limit = 600 if method == "page" else 300  # 페이지 전략은 2배까지 통째로 둔다
    assert chunks and all(len(c.text) <= limit for c in chunks), max(len(c.text) for c in chunks)
    assert {c.page for c in chunks} == {1, 2}
    assert all(text[c.start : c.end].strip() == c.text for c in chunks if c.page == 1 and method != "sentence")


def test_normalize_strategy():
    s = chunking.normalize({"method": "fixed", "size": 50, "overlap": 999})
    assert s["size"] == 100 and s["overlap"] == 50
    with pytest.raises(ValueError):
        chunking.normalize({"method": "nope"})


def test_bm25_ranking_and_filter():
    docs = ["공문서는 한글로 작성한다", "해외직구 구매대행 등록제", "한글 맞춤법과 공문서 작성 요령"]
    idx = BM25Index([10, 20, 30], [tokenize(d) for d in docs])
    hits = idx.search(tokenize("공문서 작성"), 3)
    assert {h[0] for h in hits} == {10, 30} and hits[0][1] >= hits[1][1]
    assert idx.search(tokenize("공문서 작성"), 3, allowed={30})[0][0] == 30
    assert idx.search(tokenize("없는단어"), 3) == []


def test_fusion():
    dense = [(1, 0.9), (2, 0.8), (3, 0.1)]
    bm25 = [(3, 12.0), (1, 5.0)]
    fused = rrf({"dense": dense, "bm25": bm25})
    assert fused[0]["id"] == 1 and fused[0]["ranks"] == {"dense": 1, "bm25": 2}
    w = weighted({"dense": dense, "bm25": bm25}, {"dense": 0.2, "bm25": 0.8})
    assert w[0]["id"] == 3  # BM25 가중치가 크면 BM25 1위가 올라온다


def test_rag_end_to_end(client):
    opts = client.get("/api/rag/options").json()
    assert ".hwp" in opts["file_types"] and opts["default_store"] == "qdrant"
    assert next(s for s in opts["stores"] if s["kind"] == "qdrant")["ok"] is True

    cid = client.post("/api/collections", json={"name": "통합 테스트", "domain": "legal"}).json()["id"]
    assert client.post("/api/collections", json={"name": "통합 테스트"}).status_code == 409
    files = [("files", (n, t.encode(), "text/plain")) for n, t in DOCS.items()] + [
        ("files", ("중복.txt", LAW.encode(), "text/plain")),
        ("files", ("x.exe", b"MZ", "application/octet-stream")),
    ]
    up = client.post(f"/api/collections/{cid}/documents", files=files).json()
    assert len(up["added"]) == 3 and {s["filename"] for s in up["skipped"]} == {"중복.txt", "x.exe"}

    prev = client.post(f"/api/collections/{cid}/chunk-preview", json={"method": "legal", "size": 120, "overlap": 0}).json()
    assert prev["stats"]["count"] >= 8 and any(s["heading"] == "【주문】" for s in prev["samples"])

    cs = client.post(
        f"/api/collections/{cid}/chunk-sets", json={"strategy": {"method": "legal", "size": 120, "overlap": 0, "min_chars": 5}}
    ).json()
    assert wait_job(client, cs["job"])["status"] == "done"
    chunks = client.get(f"/api/chunk-sets/{cs['id']}/chunks?limit=200").json()
    assert chunks["total"] >= 8 and all(c["tokens"] for c in chunks["items"])

    emb, rr = model_id(client, "nlpai-lab/KURE-v1"), model_id(client, "BAAI/bge-reranker-v2-m3")
    assert client.post("/api/indexes", json={"chunk_set_id": cs["id"], "model_id": rr}).status_code == 400  # 리랭커로는 인덱스 불가
    idx = client.post("/api/indexes", json={"chunk_set_id": cs["id"], "model_id": emb}).json()
    assert wait_job(client, idx["job"])["status"] == "done"
    detail = client.get(f"/api/collections/{cid}").json()
    index = detail["indexes"][0]
    assert index["status"] == "ready" and index["n_vectors"] == chunks["total"] and index["store"] == "qdrant"

    q = "해외직구 구매대행업체 등록제 시행"
    configs = [
        {"label": "의미", "index_id": index["id"], "mode": "dense", "top_k": 3},
        {"label": "BM25", "chunk_set_id": cs["id"], "mode": "bm25", "top_k": 3},
        {"label": "하이브리드+리랭커", "index_id": index["id"], "mode": "hybrid", "reranker_id": rr, "top_k": 3},
    ]
    res = client.post(f"/api/collections/{cid}/search", json={"query": q, "configs": configs}).json()["results"]
    for r in res:
        assert "error" not in r, r
        assert r["results"][0]["filename"] == "보도자료.txt", (r["config"]["label"], r["results"][0])
    assert res[2]["reranked"] and "rerank" in res[2]["results"][0]["scores"] and "before_rerank" in res[2]["results"][0]["ranks"]
    assert set(res[2]["timings_ms"]) >= {"embed", "vector_search", "bm25", "rerank"}

    # 문서 필터
    law_doc = next(d["id"] for d in detail["documents"] if d["filename"] == "법령.txt")
    filt = client.post(f"/api/collections/{cid}/search", json={"query": q, "configs": [configs[0] | {"doc_ids": [law_doc]}]}).json()[
        "results"
    ][0]
    assert all(r["filename"] == "법령.txt" for r in filt["results"])

    # 문서 삭제 → 벡터와 청크도 정리, 청크 세트는 '문서 변경됨' 표시
    press = next(d["id"] for d in detail["documents"] if d["filename"] == "보도자료.txt")
    assert client.delete(f"/api/documents/{press}").status_code == 200
    detail = client.get(f"/api/collections/{cid}").json()
    assert detail["indexes"][0]["n_vectors"] == detail["chunk_sets"][0]["n_chunks"] < chunks["total"]
    assert detail["chunk_sets"][0]["stale"] is True
    after = client.post(f"/api/collections/{cid}/search", json={"query": q, "configs": [configs[0]]}).json()["results"][0]
    assert all(r["filename"] != "보도자료.txt" for r in after["results"])

    # 정리
    assert client.delete(f"/api/indexes/{index['id']}").status_code == 200
    assert client.delete(f"/api/chunk-sets/{cs['id']}").status_code == 200
    assert client.delete(f"/api/collections/{cid}").status_code == 200
    assert all(c["id"] != cid for c in client.get("/api/collections").json())


def test_embedding_endpoints_reject_rerankers(client):
    rr = model_id(client, "BAAI/bge-reranker-v2-m3")
    r = client.post("/api/embed", json={"model_id": rr, "texts": ["x"]})
    assert r.status_code == 400 and "임베딩" in r.json()["detail"]
    assert all(m["kind"] == "reranker" for m in client.get("/api/models?kind=reranker").json())
